# Reads the calibration sheets (calibrations/calib-*.xlsx).
import ast
from pathlib import Path

from openpyxl import load_workbook

CALIB = Path(__file__).resolve().parent.parent / "1-workaround-detection-package" / "calibrations"
FILES = {"Isolation Forest": "calib-if.xlsx", "DBSCAN": "calib-dbscan.xlsx",
         "SOM": "calib-som.xlsx", "Autoencoder": "calib-auto.xlsx"}
SHORT = {"AverageWidgetSwitchingDelay": "AWSD", "EditCount": "EC", "InitialSaveTime": "IST",
         "LogLinesCount": "LLC", "UsedWidgetsCount": "UWC", "WidgetsRevisitCount": "WRC",
         "RecordStartHesitationTime": "RSHT"}


def widget_set(text):
    text = text.strip()
    if text in ("()", ""):
        return frozenset()
    names = ast.literal_eval(text)
    if isinstance(names, str):
        names = (names,)
    return frozenset(SHORT[n] for n in names)


def label(s):
    return "{" + ", ".join(sorted(s)) + "}" if s else "data only"


def load(detector):
    """{form: {widget set: dict(f, p, r, tp, k, n)}} for form in conference, employee."""
    wb = load_workbook(CALIB / FILES[detector], read_only=True, data_only=True)
    out = {}
    for form in ("conference", "employee"):
        rows = {}
        for r in wb[form].iter_rows(min_row=2, values_only=True):
            if r[0] is None:
                continue
            per_type = ast.literal_eval(r[5])
            tp = sum(v["detected_count"] for v in per_type.values())
            n = sum(v["count"] for v in per_type.values())
            f, p, rec = (float(r[0]), float(r[1]), float(r[2])) if tp else (0.0, 0.0, 0.0)
            k = round(tp / p) if tp else 0
            rows[widget_set(r[3])] = dict(f=f, p=p, r=rec, tp=tp, k=k, n=n)
        out[form] = rows
    return out


def harmonic(a, b):
    return 2 / (1 / a + 1 / b) if a > 0 and b > 0 else 0.0
