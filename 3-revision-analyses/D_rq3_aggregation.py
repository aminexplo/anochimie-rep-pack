# Ranking of the misalignment categories (RQ3) under different ways of combining effort and task success.
# Inputs are the averages in sheet "workarounds" of workarounds_assessment-v3.xlsx.
import csv
from math import sqrt
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
XLSX = HERE.parent / "1-workaround-detection-package" / "workarounds_assessment-v3.xlsx"
OUT = HERE / "out" / "D_rq3.csv"
FORMS = {"conference": (6, 9), "employee": (13, 16)}  # rows of save time ratio and data quality
METHODS = ["harmonic", "arithmetic", "geometric", "weighted_harmonic"]


def scores(e, t):
    s = 6 - t
    return {
        "harmonic": 2 / (1 / e + 1 / s),
        "arithmetic": (e + s) / 2,
        "geometric": sqrt(e * s),
        "weighted_harmonic": 1 / (0.75 / e + 0.25 / s),
    }


ws = load_workbook(XLSX, data_only=True)["workarounds"]
cats = [c.value for c in ws[1][1:6]]
rows = []
for form, (r_time, r_quality) in FORMS.items():
    time = [float(c.value) for c in ws[r_time][1:6]]
    quality = [float(c.value) for c in ws[r_quality][1:6]]
    effort = [1 + 4 * (x - min(time)) / (max(time) - min(time)) for x in time]
    success = [1 + 4 * (q - 1) / 3 for q in quality]
    table = [dict(category=c, effort=e, task_success=t, **scores(e, t)) for c, e, t in zip(cats, effort, success)]
    for m in METHODS:
        for rank, d in enumerate(sorted(table, key=lambda d: -d[m]), 1):
            d["rank_" + m] = rank
    for d in table:
        rows.append([form, d["category"], round(d["effort"], 2), round(d["task_success"], 2)]
                    + [round(d[m], 3) for m in METHODS] + [d["rank_" + m] for m in METHODS])

OUT.parent.mkdir(exist_ok=True)
with open(OUT, "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["form", "category", "effort", "task_success"] + METHODS + ["rank_" + m for m in METHODS])
    w.writerows(rows)

for r in rows:
    print(*r, sep="\t")
