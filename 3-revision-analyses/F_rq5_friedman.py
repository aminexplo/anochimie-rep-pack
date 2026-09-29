# Friedman test and Kendall's W on the developer ratings (RQ5).
# Blocks: 12 judgements (6 developers x 2 forms). Treatments: the 8 LLMs.

import csv
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.stats import friedmanchisquare, rankdata

HERE = Path(__file__).resolve().parent
XLSX = HERE.parent / "2-insight-generation-package" / "user-study-dev-results.xlsx"
OUT = HERE / "out" / "F_rq5.csv"
CRITERIA = ["Completeness", "Clarity", "Difficulty"]


def load_ratings():
    wb = load_workbook(XLSX, data_only=True)
    ratings = {c: [] for c in CRITERIA}
    for sheet in ["emp", "conf"]:
        for row in wb[sheet].iter_rows(min_row=2, values_only=True):
            if row[1] in ratings:
                ratings[row[1]].append([float(v) for v in row[2:10]])
    return {c: np.array(r) for c, r in ratings.items()}


def kendall_w(m):
    n, k = m.shape
    ranks = np.array([rankdata(r) for r in m])
    s = ((ranks.sum(axis=0) - n * (k + 1) / 2) ** 2).sum()
    ties = sum((t ** 3 - t).sum() for t in (np.unique(r, return_counts=True)[1] for r in m))
    return 12 * s / (n ** 2 * (k ** 3 - k) - n * ties)


rows = []
for crit, m in load_ratings().items():
    n, k = m.shape
    stat, p = friedmanchisquare(*m.T)
    rows.append([crit, n, k, round(stat, 3), k - 1, round(p, 4), round(kendall_w(m), 3)])

OUT.parent.mkdir(exist_ok=True)
with open(OUT, "w", newline="") as f:
    writer = csv.writer(f)
    writer.writerow(["criterion", "n", "k", "chi2", "df", "p", "kendall_w"])
    writer.writerows(rows)

for r in rows:
    print(*r, sep="\t")
