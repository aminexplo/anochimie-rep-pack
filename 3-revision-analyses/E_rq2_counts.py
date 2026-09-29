# Raw counts behind the RQ2 precision and category recall.
# One sheet per LLM in workarounds_assessment-v3.xlsx. An item is a listed workaround with a
# label (w1, w2, ...) or a judgement; lines restating an item carry neither and are not counted.
import csv
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
XLSX = HERE.parent / "1-workaround-detection-package" / "workarounds_assessment-v3.xlsx"
OUT = HERE / "out" / "E_rq2.csv"
SHEETS = {"DeepSeek-R1": "Deepseek-R1", "DeepSeek-V3": "Deepseek-V3", "Gemini-2.5-Flash": "Gemini-2.5-Flash",
          "Gemini-2.5-Pro": "Gemini-2.5-Pro", "GPT-4.5": "GPT-4.5", "GPT-4o": "GPT-4o",
          "Llama-4-Maverick": "Llama4-Maverick", "Qwen3-235B-Instruct": "Qwen3-235B"}


def blocks(ws):
    form, out = None, {}
    for r in ws.iter_rows(values_only=True):
        head = str(r[1] or "")
        if head.startswith("Predicted conference"):
            form = "conference"
        elif head.startswith("Predicted employee"):
            form = "employee"
        elif form:
            out.setdefault(form, []).append(r)
    return out


wb = load_workbook(XLSX, data_only=True)
rows = []
for model, sheet in SHEETS.items():
    for form, lines in blocks(wb[sheet]).items():
        items = [r for r in lines if r[1] and (r[0] or r[2] is not None)]
        correct = sum(1 for r in items if r[2] == 1)
        covered = sum(1 for r in lines if r[6] == 1)
        scenarios = sum(1 for r in lines if r[5])
        rows.append([model, form, len(items), correct, round(correct / len(items), 2),
                     covered, scenarios, round(covered / scenarios, 2)])

OUT.parent.mkdir(exist_ok=True)
with open(OUT, "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["model", "form", "items", "correct", "precision", "categories_covered", "categories", "recall"])
    w.writerows(rows)

for r in rows:
    print(*r, sep="\t")
